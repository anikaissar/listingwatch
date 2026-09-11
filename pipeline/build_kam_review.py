#!/usr/bin/env python3
"""Builds qa_kam_review.xlsx -- the lean, 2-sheet version of the Flipkart vs
D2C QA findings meant for a KAM (and their manager) to actually work from,
as opposed to qa_findings_flipkart_vs_d2c_v2.xlsx (the fuller 4-sheet
analyst workbook, which still exists separately -- nothing there was
removed, this is a second, narrower file).

Sheet 1 "Priority Review": one row per flagged SKU, classified into exactly
one priority tier (P1 most urgent -> P4 least), sorted, with a short
"Status" (Open / Resolved / Not an Issue) and "Reminders Sent" column the
KAM edits directly in Excel. Both columns are LIVE cells, not formulas --
Excel is the source of truth for them between rebuilds of this report.

Sheet 2 "Summary": a manager-readable rollup -- total / open / resolved per
tier, and total reminders sent -- as LIVE FORMULAS referencing Sheet 1, so
it updates automatically the moment someone edits Status or Reminders Sent
in Excel. No need to come back to Claude for that to work.

IMPORTANT -- what persisting Status/Reminders Sent across a data refresh
actually requires: this script has no live connection to whatever copy of
the file is sitting on the KAM's machine. The only way it can know what
they've marked is if that edited file is handed back to it (as --existing)
before it rebuilds. If you skip that step, a rebuild starts every row over
at Open/0. Given that file, it carries every (nh_sku, fsn)'s Status forward
onto the freshly computed rows, and a row silently drops off the sheet only
when the fresh diff data no longer flags it under any tier at all (a manual
Status never causes a still-flagged row to disappear).

Reminders Sent is no longer a purely manual field -- as of 2026-09-04 it is
auto-incremented on rebuild, per row, as follows: a row seen for the first
time starts at 0 (nothing to remind anyone of yet). A row that was already
Open and is still flagged gets +1 (another cycle passed with no action, so
that's a genuine additional reminder). A row marked "Not an Issue" keeps its
count frozen forever, however many more times it's rebuilt -- that status
is a deliberate KAM judgment call to stop chasing it, and the point of this
field is to never let that judgment quietly rack up reminders that make the
KAM or their manager look unresponsive. A row marked "Resolved" that is
STILL flagged by fresh data is reopened back to "Open" (with a note appended
to its detail) and does get +1 -- because that combination means the fix
didn't actually take, which is worth surfacing, not hiding under a
Resolved label.

Usage:
    python build_kam_review.py [--diff qa_diff_flipkart_vs_d2c_v4.csv] [--existing qa_kam_review.xlsx] [--out qa_kam_review.xlsx]
"""
import argparse
import csv
import datetime
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
NOTE_FONT = Font(name=FONT, italic=True, size=10, color="555555")
BOLD = Font(name=FONT, bold=True)
TITLE_FONT = Font(name=FONT, bold=True, size=14)
SECTION_FONT = Font(name=FONT, bold=True, size=12, color="1F4E5F")
RESOLVED_FONT = Font(name=FONT, color="8A8A8A", strike=True)
TIER_FILLS = {
    "P1": PatternFill("solid", fgColor="E06666"),
    "P2": PatternFill("solid", fgColor="F6B26B"),
    "P3": PatternFill("solid", fgColor="FFD966"),
    "P4": PatternFill("solid", fgColor="CFE2F3"),
}
STATUS_OPTIONS = ["Open", "Resolved", "Not an Issue"]

# 4-tier scheme (as of 2026-09-08). The old scheme had 5 tiers, with P1 (Flipkart's
# flat "3 Months" default) and P2 (any other shelf-life mismatch) as separate tiers;
# those are now one tier, P1, and everything below shifted up a number: old P3 (broken
# listing) -> P2, old P4 (title check) -> P3, old P5 (photo check) -> P4. The two old
# P1/P2 sub-cases are still distinguishable within the new P1 via the detail text --
# see classify() below.
TIER_LABEL = {
    "P1": "Shelf Life Mismatch",
    "P2": "Page Broken",
    "P3": "Title Mismatch",
    "P4": "Product Photo Mismatch",
}
# Short line for the top-of-sheet action plan -- kept to one clause each.
TIER_ACTION = {
    "P1": "Escalate flat 3-month defaults once; fix other mismatches per SKU",
    "P2": "Check listing status (dead FSN? soft-blocked?)",
    "P3": "Confirm this is really the right SKU",
    "P4": "Spot-check the product photos",
}

_DURATION_RE = re.compile(r"(\d+)\s*(day|days|month|months|year|years)", re.I)


def _short_duration(text):
    m = _DURATION_RE.search(text or "")
    if not m:
        return (text or "").strip()[:20] or "?"
    n, u = m.group(1), m.group(2).lower().rstrip("s")
    return f"{n} {u}{'s' if n != '1' else ''}"


def classify(rr):
    """Same tiering as build_report.py's Priority Review, kept independent
    here so this script stays self-contained. Returns (tier, issue, detail)
    or (None, None, None) if the row isn't flagged at all.

    P1 (Shelf Life Mismatch) covers two sub-cases that used to be separate
    tiers: Flipkart's flat "3 Months" default (one escalation to Flipkart's
    catalog team covers every SKU showing this) and any other shelf-life
    mismatch (each needs its own per-SKU fix). Both still say "P1" here, but
    the detail text keeps them tellable apart -- the flat-default rows say
    so explicitly."""
    if rr["shelf_life_status"] == "mismatch":
        fk_val = (rr.get("flipkart_max_shelf_life") or "").strip()
        d2c_short = _short_duration(rr.get("d2c_shelf_life", ""))
        if fk_val == "3 Months":
            return "P1", TIER_LABEL["P1"], f"D2C {d2c_short} vs Flipkart 3 months (flat default -- escalate once, not per SKU)"
        fk_short = _short_duration(fk_val) if fk_val else "not on Specifications grid"
        return "P1", TIER_LABEL["P1"], f"D2C {d2c_short} vs Flipkart {fk_short}"
    if rr["flipkart_page_broken"] is True:
        return "P2", TIER_LABEL["P2"], "Generic placeholder page, not real product data"
    if rr["title_plausible"] is False:
        return "P3", TIER_LABEL["P3"], "Flipkart title doesn't share the SKU's core words"
    vb = rr["visual_best_similarity"]
    if vb is not None and vb < 0.6:
        return "P4", TIER_LABEL["P4"], f"Best photo match scores {vb:.2f} (< 0.6)"
    return None, None, None


def classify_nykaa(rr):
    """Nykaa counterpart to classify() above. Same P1-P4 tier numbering and
    labels (TIER_LABEL), but P1 (Shelf Life Mismatch) never fires here --
    confirmed (2026-09-11, 45/45 real Nykaa products sampled) that Nykaa's
    own structured expiry field is null on every product checked, so there
    is nothing on Nykaa's side to diff a shelf-life claim against. Revisit
    this if Nykaa ever starts populating that field for some SKUs.

    P2 (Page Broken) covers Nykaa's own 404/isNotFound signal as well as a
    total extraction failure (no_content) -- see nykaa_qa_diff.py's
    is_nykaa_page_broken() and nykaa_pdp_scraper.py's page-not-found
    finding for the real case (product 10346740) this was built against."""
    if rr["nykaa_page_broken"] is True:
        return "P2", TIER_LABEL["P2"], "Nykaa page broken, removed, or returned no usable content"
    if rr["title_plausible"] is False:
        return "P3", TIER_LABEL["P3"], "Nykaa title doesn't share the SKU's core words"
    vb = rr["visual_best_similarity"]
    if vb is not None and vb < 0.6:
        return "P4", TIER_LABEL["P4"], f"Best photo match scores {vb:.2f} (< 0.6)"
    return None, None, None


DIFF_COLS_TYPES = {
    "flipkart_page_broken": "bool", "title_plausible": "bool",
    "visual_best_similarity": "float",
}

NYKAA_DIFF_COLS_TYPES = {
    "nykaa_page_broken": "bool", "title_plausible": "bool",
    "visual_best_similarity": "float",
}


def load_diff_rows(path, col_types=None):
    """col_types defaults to DIFF_COLS_TYPES (Flipkart's diff column set) --
    pass NYKAA_DIFF_COLS_TYPES when loading a nykaa_qa_diff.py output file,
    since its bool column is named nykaa_page_broken rather than
    flipkart_page_broken."""
    if col_types is None:
        col_types = DIFF_COLS_TYPES
    rows = []
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            for c, kind in col_types.items():
                v = r.get(c, "")
                if kind == "bool":
                    r[c] = True if v == "True" else (False if v == "False" else None)
                elif kind == "float":
                    r[c] = float(v) if v not in ("", None) else None
            rows.append(r)
    return rows


def load_prior_kam_edits(path):
    """Reads a previously-delivered qa_kam_review.xlsx's Priority Review
    sheet and returns {(nh_sku, fsn): {"status": ..., "reminders_sent": ...}}.
    Returns {} if the file doesn't exist or doesn't have the sheet/columns
    yet (e.g. the very first build) -- never raises."""
    try:
        wb = load_workbook(path, data_only=True)
    except (FileNotFoundError, Exception):  # noqa: BLE001
        return {}
    if "Priority Review" not in wb.sheetnames:
        return {}
    ws = wb["Priority Review"]
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=min(30, ws.max_row)):
        values = [c.value for c in row]
        if values and values[0] == "priority":
            header_row = row[0].row
            break
    if header_row is None:
        return {}
    headers = [c.value for c in ws[header_row]]
    try:
        i_status = headers.index("status")
        i_sku = headers.index("nh_sku")
        i_fsn = headers.index("fsn")
        i_rem = headers.index("reminders_sent")
    except ValueError:
        return {}
    edits = {}
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        sku, fsn = row[i_sku].value, row[i_fsn].value
        if not sku or not fsn:
            continue
        status = row[i_status].value or "Open"
        reminders = row[i_rem].value or 0
        edits[(sku, fsn)] = {"status": status, "reminders_sent": reminders}
    return edits


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 22


def build(diff_path, existing_path, out_path):
    diff_rows = load_diff_rows(diff_path)
    prior_edits = load_prior_kam_edits(existing_path)

    classified = []
    recurred = 0
    for rr in diff_rows:
        tier, issue, detail = classify(rr)
        if tier:
            key = (rr["nh_sku"], rr["fsn"])
            prior = prior_edits.get(key)
            if prior is None:
                status, reminders = "Open", 0
            elif prior["status"] == "Not an Issue":
                status, reminders = "Not an Issue", prior["reminders_sent"]
            elif prior["status"] == "Resolved":
                status, reminders = "Open", prior["reminders_sent"] + 1
                detail = detail + " [recurred after being marked Resolved]"
                recurred += 1
            else:  # still Open
                status, reminders = "Open", prior["reminders_sent"] + 1
            classified.append((tier, rr, issue, detail, status, reminders))
    classified.sort(key=lambda t: (t[0], t[1]["nh_sku"]))

    carried = sum(1 for *_, status, rem in classified if status != "Open" or rem)
    dropped = len(prior_edits) - sum(
        1 for tier, rr, *_ in classified if (rr["nh_sku"], rr["fsn"]) in prior_edits
    )

    wb = Workbook()

    # ---------------- Sheet 1: Priority Review ----------------
    ws1 = wb.active
    ws1.title = "Priority Review"
    ws1["A1"] = "Priority Review — work top to bottom"
    ws1["A1"].font = TITLE_FONT
    ws1["A2"] = ("Generated " + datetime.date.today().isoformat() +
                 ". Status is yours to edit directly in Excel — the Summary sheet "
                 "updates live. Reminders Sent auto-increments each week a row stays "
                 "Open; mark a row Not an Issue to freeze its count for good, or "
                 "Resolved once it's actually fixed (it reopens itself if Flipkart "
                 "still shows the problem next time).")
    ws1["A2"].font = NOTE_FONT
    ws1.merge_cells("A2:F2")
    ws1["A2"].alignment = Alignment(wrap_text=True)

    r = 4
    ws1.cell(row=r, column=1, value="Order of work:").font = SECTION_FONT
    r += 1
    tier_counts = {t: sum(1 for c in classified if c[0] == t) for t in TIER_ACTION}
    for tier in ["P1", "P2", "P3", "P4"]:
        cnt = tier_counts.get(tier, 0)
        c = ws1.cell(row=r, column=1, value=tier)
        c.font = Font(name=FONT, bold=True)
        c.fill = TIER_FILLS[tier]
        ws1.cell(row=r, column=2, value=f"{cnt}").font = BOLD
        ws1.cell(row=r, column=3, value=TIER_ACTION[tier]).font = Font(name=FONT)
        r += 1
    r += 1

    header_row = r
    headers = ["priority", "status", "nh_sku", "fsn", "issue", "detail",
               "flipkart_title", "d2c_title", "reminders_sent"]
    for ci, h in enumerate(headers, start=1):
        ws1.cell(row=header_row, column=ci, value=h)
    style_header(ws1, len(headers), row=header_row)

    dv = DataValidation(type="list", formula1='"Open,Resolved,Not an Issue"', allow_blank=False)
    ws1.add_data_validation(dv)

    for tier, rr, issue, detail, status, reminders in classified:
        ws1.append([tier, status, rr["nh_sku"], rr["fsn"], issue, detail,
                    rr["flipkart_title"], rr["d2c_title"], reminders])
    data_first_row = header_row + 1
    data_last_row = ws1.max_row
    dv.add(f"B{data_first_row}:B{data_last_row}")

    for row in ws1.iter_rows(min_row=data_first_row, max_row=data_last_row):
        status_val = row[1].value
        font = RESOLVED_FONT if status_val in ("Resolved", "Not an Issue") else Font(name=FONT)
        for cell in row:
            cell.font = font

    for tier, fill in TIER_FILLS.items():
        ws1.conditional_formatting.add(
            f"A{data_first_row}:A{data_last_row}",
            CellIsRule(operator="equal", formula=[f'"{tier}"'], fill=fill),
        )

    ws1.freeze_panes = f"A{data_first_row}"
    ws1.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{data_last_row}"
    autosize(ws1, [10, 14, 16, 18, 28, 34, 40, 40, 14])

    # ---------------- Sheet 2: Summary ----------------
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Flipkart vs D2C Listing QA — Summary"
    ws2["A1"].font = TITLE_FONT
    ws2["A2"] = (f"Generated {datetime.date.today().isoformat()}. Totals below are live "
                 f"formulas reading the 'Priority Review' sheet — they update the moment "
                 f"Status or Reminders Sent changes there, no rebuild needed.")
    ws2["A2"].font = NOTE_FONT
    ws2.merge_cells("A2:F2")
    ws2["A2"].alignment = Alignment(wrap_text=True)

    pr = "'Priority Review'!"
    prio_rng = f"{pr}A{data_first_row}:A{data_last_row}"
    status_rng = f"{pr}B{data_first_row}:B{data_last_row}"
    rem_rng = f"{pr}I{data_first_row}:I{data_last_row}"

    r = 4
    headers2 = ["Priority", "What it means", "Total", "Open", "Resolved", "Not an Issue", "Reminders Sent"]
    for ci, h in enumerate(headers2, start=1):
        ws2.cell(row=r, column=ci, value=h)
    style_header(ws2, len(headers2), row=r)
    r += 1
    first_data_row = r
    for tier in ["P1", "P2", "P3", "P4"]:
        c1 = ws2.cell(row=r, column=1, value=tier)
        c1.font = Font(name=FONT, bold=True)
        c1.fill = TIER_FILLS[tier]
        ws2.cell(row=r, column=2, value=TIER_LABEL[tier]).font = Font(name=FONT)
        ws2.cell(row=r, column=3, value=f'=COUNTIF({prio_rng},"{tier}")').font = BOLD
        ws2.cell(row=r, column=4,
                 value=f'=COUNTIFS({prio_rng},"{tier}",{status_rng},"Open")').font = Font(name=FONT)
        ws2.cell(row=r, column=5,
                 value=f'=COUNTIFS({prio_rng},"{tier}",{status_rng},"Resolved")').font = Font(name=FONT)
        ws2.cell(row=r, column=6,
                 value=f'=COUNTIFS({prio_rng},"{tier}",{status_rng},"Not an Issue")').font = Font(name=FONT)
        ws2.cell(row=r, column=7,
                 value=f'=SUMIFS({rem_rng},{prio_rng},"{tier}")').font = Font(name=FONT)
        r += 1
    last_data_row = r - 1

    total_row = r
    ws2.cell(row=r, column=1, value="Total").font = BOLD
    ws2.cell(row=r, column=3, value=f"=SUM(C{first_data_row}:C{last_data_row})").font = BOLD
    ws2.cell(row=r, column=4, value=f"=SUM(D{first_data_row}:D{last_data_row})").font = BOLD
    ws2.cell(row=r, column=5, value=f"=SUM(E{first_data_row}:E{last_data_row})").font = BOLD
    ws2.cell(row=r, column=6, value=f"=SUM(F{first_data_row}:F{last_data_row})").font = BOLD
    ws2.cell(row=r, column=7, value=f"=SUM(G{first_data_row}:G{last_data_row})").font = BOLD
    r += 2

    ws2.cell(row=r, column=1, value="Note").font = SECTION_FONT
    r += 1
    notes = [
        "P1 (Shelf Life Mismatch) mixes two kinds of fix: rows whose Detail says \"flat default\" are Flipkart's blanket \"3 Months\" value — one escalation to Flipkart's catalog team covers all of those, not a per-SKU fix. Every other P1 row is a genuine per-SKU shelf-life mismatch that needs its own fix.",
        "Status is edited directly in this file; Reminders Sent auto-increments each week a row stays Open, freezes for good once marked Not an Issue, and reopens itself (with a note) if something marked Resolved is still flagged. A row only disappears on its own once a fresh data pull confirms the underlying issue is actually gone — marking it here doesn't remove it, by design, so nothing gets silently lost.",
        "Full detail (raw per-listing data, image/description similarity, the missing-D2C-reference list) lives in the separate qa_findings_flipkart_vs_d2c_v2.xlsx analyst workbook — this file is deliberately just the actionable list and its rollup.",
    ]
    for note in notes:
        ws2.cell(row=r, column=1, value="• " + note).font = Font(name=FONT, size=10)
        ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws2.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws2.row_dimensions[r].height = 30
        r += 1

    autosize(ws2, [10, 34, 8, 8, 10, 12, 14])

    wb.save(out_path)
    print(f"Saved {out_path} — Priority Review: {len(classified)} rows "
          f"({tier_counts}); carried forward {carried} prior edit(s) from {existing_path}"
          + (f"; {dropped} previously-tracked row(s) no longer flagged and dropped off" if dropped > 0 else "")
          + (f"; {recurred} row(s) reopened -- still flagged despite being marked Resolved" if recurred > 0 else ""))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--diff", default="qa_diff_flipkart_vs_d2c_v4.csv")
    ap.add_argument("--existing", default="qa_kam_review.xlsx",
                     help="a previously-delivered copy of this same file, to carry forward Status/Reminders Sent")
    ap.add_argument("--out", default="qa_kam_review.xlsx")
    a = ap.parse_args()
    build(a.diff, a.existing, a.out)
